from __future__ import annotations

import argparse
from datetime import date
import math
import re
import sys
from pathlib import Path
from typing import Iterable

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

SKILL_SCRIPTS = (
    Path(r"C:\Users\Admin\.codex\plugins\cache\openai-primary-runtime\documents")
    / "26.619.11828"
    / "skills"
    / "documents"
    / "scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


TITLE = "Bảo mật hệ thống RESTful API cho dịch vụ khóa học online nhỏ"
SUBTITLE = (
    "Sử dụng Spring Boot, NextJS, JWT, bcrypt, AES-GCM, HMAC-SHA256, HTTPS/TLS "
    "và Swagger/OpenAPI"
)

REPORT_PRESET = {
    "preset_name": "narrative_proposal + academic_times_override",
    "base_font": "Times New Roman",
    "body_size": 13,
    "body_after": 6,
    "body_line": 1.3,
    "heading1": {"size": 17, "color": "1F3A5F", "before": 18, "after": 10},
    "heading2": {"size": 14, "color": "1F3A5F", "before": 12, "after": 6},
    "heading3": {"size": 13, "color": "202124", "before": 8, "after": 4},
    "table_header_fill": "EFF3F8",
    "table_border": "C9D2DD",
}
EXPORT_DATE_TEXT = date.today().strftime("%d tháng %m năm %Y")


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    name: str = REPORT_PRESET["base_font"],
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = REPORT_PRESET["body_after"],
    line_spacing: float = REPORT_PRESET["body_line"],
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    paragraph.alignment = alignment


def add_shading(element, fill: str):
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=9.5, color="6F6F6F")


def set_page_number_start(section, start: int):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def ensure_style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    styles = doc.styles
    if name in styles:
        return styles[name]
    return styles.add_style(name, style_type)


def configure_section(section, *, running_header_footer: bool):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False

    if not running_header_footer:
        return

    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=0, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(TITLE)
    set_run_font(run, size=9.5, color="6F6F6F", italic=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=0, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    label = p.add_run("Trang ")
    set_run_font(label, size=9.5, color="6F6F6F")
    add_page_number(p)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = REPORT_PRESET["base_font"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    normal.font.size = Pt(REPORT_PRESET["body_size"])
    normal.font.color.rgb = rgb("202124")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(REPORT_PRESET["body_after"])
    normal.paragraph_format.line_spacing = REPORT_PRESET["body_line"]

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = REPORT_PRESET["base_font"]
    heading1._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    heading1.font.size = Pt(REPORT_PRESET["heading1"]["size"])
    heading1.font.color.rgb = rgb(REPORT_PRESET["heading1"]["color"])
    heading1.font.bold = True
    heading1.paragraph_format.space_before = Pt(REPORT_PRESET["heading1"]["before"])
    heading1.paragraph_format.space_after = Pt(REPORT_PRESET["heading1"]["after"])
    heading1.paragraph_format.line_spacing = 1.1

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = REPORT_PRESET["base_font"]
    heading2._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    heading2._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    heading2.font.size = Pt(REPORT_PRESET["heading2"]["size"])
    heading2.font.color.rgb = rgb(REPORT_PRESET["heading2"]["color"])
    heading2.font.bold = True
    heading2.paragraph_format.space_before = Pt(REPORT_PRESET["heading2"]["before"])
    heading2.paragraph_format.space_after = Pt(REPORT_PRESET["heading2"]["after"])
    heading2.paragraph_format.line_spacing = 1.08

    heading3 = doc.styles["Heading 3"]
    heading3.font.name = REPORT_PRESET["base_font"]
    heading3._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    heading3._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    heading3._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    heading3.font.size = Pt(REPORT_PRESET["heading3"]["size"])
    heading3.font.color.rgb = rgb(REPORT_PRESET["heading3"]["color"])
    heading3.font.bold = True
    heading3.paragraph_format.space_before = Pt(REPORT_PRESET["heading3"]["before"])
    heading3.paragraph_format.space_after = Pt(REPORT_PRESET["heading3"]["after"])
    heading3.paragraph_format.line_spacing = 1.05

    cover_kicker = ensure_style(doc, "CoverKicker")
    cover_kicker.font.name = REPORT_PRESET["base_font"]
    cover_kicker._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    cover_kicker._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    cover_kicker._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    cover_kicker.font.size = Pt(12)
    cover_kicker.font.bold = True
    cover_kicker.font.color.rgb = rgb("55606D")
    cover_kicker.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_kicker.paragraph_format.space_before = Pt(0)
    cover_kicker.paragraph_format.space_after = Pt(8)
    cover_kicker.paragraph_format.line_spacing = 1.0

    cover_title = ensure_style(doc, "CoverTitle")
    cover_title.font.name = REPORT_PRESET["base_font"]
    cover_title._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    cover_title._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    cover_title._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    cover_title.font.size = Pt(25)
    cover_title.font.bold = True
    cover_title.font.color.rgb = rgb("163B64")
    cover_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_before = Pt(0)
    cover_title.paragraph_format.space_after = Pt(8)
    cover_title.paragraph_format.line_spacing = 1.08

    cover_subtitle = ensure_style(doc, "CoverSubtitle")
    cover_subtitle.font.name = REPORT_PRESET["base_font"]
    cover_subtitle._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    cover_subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    cover_subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    cover_subtitle.font.size = Pt(13)
    cover_subtitle.font.color.rgb = rgb("4B5563")
    cover_subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_subtitle.paragraph_format.space_before = Pt(0)
    cover_subtitle.paragraph_format.space_after = Pt(16)
    cover_subtitle.paragraph_format.line_spacing = 1.15

    meta_style = ensure_style(doc, "CoverMeta")
    meta_style.font.name = REPORT_PRESET["base_font"]
    meta_style._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    meta_style._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    meta_style._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    meta_style.font.size = Pt(11.5)
    meta_style.font.color.rgb = rgb("202124")
    meta_style.paragraph_format.space_before = Pt(0)
    meta_style.paragraph_format.space_after = Pt(4)
    meta_style.paragraph_format.line_spacing = 1.15

    toc_title = ensure_style(doc, "TOCTitle")
    toc_title.font.name = REPORT_PRESET["base_font"]
    toc_title._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    toc_title._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    toc_title._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    toc_title.font.size = Pt(18)
    toc_title.font.bold = True
    toc_title.font.color.rgb = rgb("163B64")
    toc_title.paragraph_format.space_before = Pt(0)
    toc_title.paragraph_format.space_after = Pt(12)
    toc_title.paragraph_format.line_spacing = 1.0

    toc_item = ensure_style(doc, "TOCItem")
    toc_item.font.name = REPORT_PRESET["base_font"]
    toc_item._element.rPr.rFonts.set(qn("w:ascii"), REPORT_PRESET["base_font"])
    toc_item._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_PRESET["base_font"])
    toc_item._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_PRESET["base_font"])
    toc_item.font.size = Pt(11.5)
    toc_item.font.color.rgb = rgb("202124")
    toc_item.paragraph_format.space_before = Pt(0)
    toc_item.paragraph_format.space_after = Pt(5)
    toc_item.paragraph_format.line_spacing = 1.15

    code_style = ensure_style(doc, "CodeBlock")
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.font.size = Pt(9.5)
    code_style.font.color.rgb = rgb("1F2937")
    code_style.paragraph_format.space_before = Pt(2)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.1
    code_style.paragraph_format.left_indent = Inches(0.22)
    code_style.paragraph_format.right_indent = Inches(0.08)


def add_cover(doc: Document):
    for _ in range(6):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=0, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run("")

    p = doc.add_paragraph(style="CoverKicker")
    p.add_run("BÁO CÁO ĐỒ ÁN MÔN MẬT MÃ ỨNG DỤNG")

    p = doc.add_paragraph(style="CoverTitle")
    p.add_run(TITLE)

    p = doc.add_paragraph(style="CoverSubtitle")
    p.add_run(SUBTITLE)

    p = doc.add_paragraph(style="CoverSubtitle")
    p.add_run("Phiên bản hoàn chỉnh: Tóm tắt, Lời mở đầu, Chương 1 đến Chương 5, Tài liệu tham khảo và Phụ lục")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    labels = [
        ("Loại tài liệu", "Bản hoàn chỉnh phục vụ in, rà soát và nộp báo cáo"),
        ("Định hướng đề tài", "Mật mã ứng dụng và bảo mật RESTful API"),
        ("Công nghệ trọng tâm", "JWT, bcrypt, AES-GCM, HMAC-SHA256, HTTPS/TLS, Swagger/OpenAPI"),
        ("Ngày xuất tài liệu", EXPORT_DATE_TEXT),
    ]
    widths = [2700, 6660]
    for row_index, (label, value) in enumerate(labels):
        row = table.rows[row_index]
        row.cells[0].text = label
        row.cells[1].text = value
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            set_paragraph_spacing(para, before=0, after=0, line_spacing=1.1, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            if cell_index == 0:
                para.runs[0].bold = True
                set_run_font(para.runs[0], size=11, color="163B64", bold=True)
            else:
                set_run_font(para.runs[0], size=11.5, color="202124")
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=120)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=11.5)
    for cell in table.columns[0].cells:
        add_shading(cell._tc, "F8FAFC")

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=16, after=0, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run("Tài liệu được biên soạn từ source code và tài liệu báo cáo hiện tại của project.")
    set_run_font(run, size=10.5, color="6B7280", italic=True)


def add_toc_page(doc: Document, chapter_titles: Iterable[str]):
    p = doc.add_paragraph(style="TOCTitle")
    p.add_run("MỤC LỤC")

    intro = doc.add_paragraph(style="TOCItem")
    set_paragraph_spacing(intro, before=0, after=10, line_spacing=1.15, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    intro.add_run("Tài liệu này trình bày bản báo cáo hoàn chỉnh theo hướng đồ án môn Mật mã ứng dụng và bảo mật API.")

    for level, title in chapter_titles:
        p = doc.add_paragraph(style="TOCItem")
        set_paragraph_spacing(p, before=0, after=6, line_spacing=1.15, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Inches(0.15 + (0.28 * level))
        run = p.add_run(title)
        set_run_font(
            run,
            size=11.5 if level == 0 else 11,
            color="202124",
            bold=level == 0,
        )


def extract_toc_titles(markdown_text: str) -> list[tuple[int, str]]:
    titles: list[tuple[int, str]] = []
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            titles.append((0, stripped.replace("## ", "", 1).strip()))
        elif stripped.startswith("### "):
            titles.append((1, stripped.replace("### ", "", 1).strip()))
    return titles


def render_inline(paragraph, node):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=REPORT_PRESET["body_size"], color="202124")
        return

    if not isinstance(node, Tag):
        return

    if node.name == "br":
        paragraph.add_run().add_break()
        return

    if node.name in {"strong", "b"}:
        text = node.get_text()
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=REPORT_PRESET["body_size"], color="202124", bold=True)
        return

    if node.name in {"em", "i"}:
        text = node.get_text()
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=REPORT_PRESET["body_size"], color="202124", italic=True)
        return

    if node.name == "code":
        text = node.get_text()
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, name="Consolas", size=9.5, color="163B64")
        return

    for child in node.children:
        render_inline(paragraph, child)


def paragraph_text_length(paragraphs: list[str]) -> int:
    value = max((len(p.strip()) for p in paragraphs if p.strip()), default=8)
    return max(8, value)


def add_html_table(doc: Document, element: Tag):
    rows = []
    for tr in element.find_all("tr"):
        cells = tr.find_all(["th", "td"])
        rows.append([cell.get_text(" ", strip=True) for cell in cells])
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    max_lens = []
    for column_index in range(column_count):
        lengths = [len(row[column_index]) for row in normalized_rows]
        lens = max(lengths) if lengths else 10
        weight = max(1.0, min(4.0, 0.85 + math.sqrt(max(lens, 1)) / 3.2))
        max_lens.append(weight)
    widths = column_widths_from_weights(max_lens)

    table = doc.add_table(rows=len(normalized_rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for row_index, row_data in enumerate(normalized_rows):
        row = table.rows[row_index]
        for cell_index, value in enumerate(row_data):
            cell = row.cells[cell_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 24 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(para, before=0, after=0, line_spacing=1.1, alignment=alignment)
            run = para.add_run(value)
            if row_index == 0:
                set_run_font(run, size=11, color="163B64", bold=True)
            else:
                set_run_font(run, size=11, color="202124")
            if row_index == 0:
                add_shading(cell._tc, REPORT_PRESET["table_header_fill"])

    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=120)

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = tc_pr.find(qn("w:tcBorders"))
            if tc_borders is None:
                tc_borders = OxmlElement("w:tcBorders")
                tc_pr.append(tc_borders)
            for side in ("top", "left", "bottom", "right"):
                border = tc_borders.find(qn(f"w:{side}"))
                if border is None:
                    border = OxmlElement(f"w:{side}")
                    tc_borders.append(border)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "6")
                border.set(qn("w:space"), "0")
                border.set(qn("w:color"), REPORT_PRESET["table_border"])

    after = doc.add_paragraph()
    set_paragraph_spacing(after, before=2, after=2, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.LEFT)


def add_code_block(doc: Document, text: str):
    p = doc.add_paragraph(style="CodeBlock")
    set_paragraph_spacing(p, before=2, after=8, line_spacing=1.08, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    add_shading(p._p, "F5F7FA")
    for index, line in enumerate(text.rstrip().splitlines()):
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=9.5, color="1F2937")
        if index != len(text.rstrip().splitlines()) - 1:
            run.add_break()


def add_list_item(doc: Document, item: Tag, *, numbered: bool, number_text: str | None = None):
    style_name = "List Bullet" if not numbered else None
    p = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    set_paragraph_spacing(
        p,
        before=0,
        after=4,
        line_spacing=1.208,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    if numbered and number_text:
        prefix = p.add_run(f"{number_text} ")
        set_run_font(prefix, size=11, color="202124")
    for child in item.children:
        if isinstance(child, Tag) and child.name in {"ul", "ol"}:
            continue
        render_inline(p, child)

    nested_lists = [child for child in item.children if isinstance(child, Tag) and child.name in {"ul", "ol"}]
    for nested in nested_lists:
        for nested_item in nested.find_all("li", recursive=False):
            nested_paragraph = doc.add_paragraph(style="List Bullet" if nested.name == "ul" else None)
            set_paragraph_spacing(
                nested_paragraph,
                before=0,
                after=3,
                line_spacing=1.15,
                alignment=WD_ALIGN_PARAGRAPH.LEFT,
            )
            nested_paragraph.paragraph_format.left_indent = Inches(0.65)
            nested_paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            if nested.name == "ol":
                nested_start = int(nested.get("start", "1"))
                nested_index = list(nested.find_all("li", recursive=False)).index(nested_item)
                prefix = nested_paragraph.add_run(f"{nested_start + nested_index}. ")
                set_run_font(prefix, size=11, color="202124")
            for child in nested_item.children:
                render_inline(nested_paragraph, child)


def add_heading(doc: Document, level: int, text: str, *, insert_break: bool):
    if insert_break:
        doc.add_page_break()
    style_name = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    p.add_run(text.strip())


def build_document(markdown_path: Path, output_path: Path):
    markdown_text = markdown_path.read_text(encoding="utf-8")
    markdown_text = re.sub(r"(?m)^(\d+)\.(\S)", r"\1. \2", markdown_text)
    toc_titles = extract_toc_titles(markdown_text)

    html = markdown.markdown(
        markdown_text,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="html5",
    )
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Báo cáo đồ án Mật mã ứng dụng"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = "Báo cáo hoàn chỉnh phục vụ nộp môn Mật mã ứng dụng"

    configure_styles(doc)
    configure_section(doc.sections[0], running_header_footer=False)
    add_cover(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section, running_header_footer=True)
    set_page_number_start(body_section, 1)
    add_toc_page(doc, toc_titles)
    doc.add_page_break()

    chapter_index = -1
    for element in soup.children:
        if not isinstance(element, Tag):
            continue

        if element.name == "h1":
            continue

        if element.name in {"h2", "h3", "h4"}:
            text = element.get_text(" ", strip=True)
            if not text:
                continue
            if element.name == "h2":
                chapter_index += 1
                add_heading(doc, 1, text, insert_break=chapter_index > 0)
            elif element.name == "h3":
                add_heading(doc, 2, text, insert_break=False)
            else:
                add_heading(doc, 3, text, insert_break=False)
            continue

        if element.name == "p":
            text = element.get_text(" ", strip=False).strip()
            if not text:
                continue
            p = doc.add_paragraph()
            alignment = WD_ALIGN_PARAGRAPH.LEFT if len(text) < 180 or "`" in text else WD_ALIGN_PARAGRAPH.JUSTIFY
            set_paragraph_spacing(
                p,
                before=0,
                after=REPORT_PRESET["body_after"],
                line_spacing=REPORT_PRESET["body_line"],
                alignment=alignment,
            )
            for child in element.children:
                render_inline(p, child)
            continue

        if element.name == "pre":
            code = element.get_text("\n", strip=False)
            add_code_block(doc, code)
            continue

        if element.name == "ul":
            for item in element.find_all("li", recursive=False):
                add_list_item(doc, item, numbered=False)
            continue

        if element.name == "ol":
            start = int(element.get("start", "1"))
            for offset, item in enumerate(element.find_all("li", recursive=False)):
                add_list_item(doc, item, numbered=True, number_text=f"{start + offset}.")
            continue

        if element.name == "table":
            add_html_table(doc, element)
            continue

        if element.name == "blockquote":
            paragraphs = [p.get_text(" ", strip=True) for p in element.find_all("p")]
            text = "\n".join(paragraphs).strip()
            if text:
                p = doc.add_paragraph()
                set_paragraph_spacing(p, before=2, after=8, line_spacing=1.15, alignment=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.15)
                add_shading(p._p, "F8FAFC")
                run = p.add_run(text)
                set_run_font(run, size=10.5, color="374151", italic=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build a clean DOCX report from the markdown draft.")
    parser.add_argument("markdown_path", type=Path)
    parser.add_argument("output_path", type=Path)
    args = parser.parse_args()
    build_document(args.markdown_path, args.output_path)


if __name__ == "__main__":
    main()
